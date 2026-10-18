# -*- coding: utf-8 -*-
"""dokumente.py — hochgeladenen Lebenslauf lesen (Word/PDF) und Ergebnisse als Word schreiben."""

import io
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

MIN_FOTO = 5_000          # kleinere Bilder sind Symbole oder Logos, kein Foto
MAX_BILD_BYTES = 15_000_000
MAX_PIXEL = 25_000_000    # Schutz gegen praeparierte Riesenbilder
FOTO_PX = 800             # das Foto wird hoechstens so gross neu gespeichert


def _bildart(daten):
    if daten[:3] == b"\xff\xd8\xff":
        return "jpg"
    if daten[:8] == b"\x89PNG\r\n\x1a\n":
        return "png"
    return None


def _sauber(daten):
    """Bild pruefen, verkleinern und neu als JPEG speichern — ohne EXIF-Angaben
    (z. B. GPS-Ort vom Handyfoto). None, wenn es kein brauchbares Bild ist."""
    try:
        from PIL import Image
        Image.MAX_IMAGE_PIXELS = MAX_PIXEL
        with Image.open(io.BytesIO(daten)) as bild:
            if bild.width * bild.height > MAX_PIXEL:
                return None
            # Ein Passfoto ist etwa so hoch wie breit; Banner und Hintergruende nicht.
            if not 0.5 <= bild.width / max(1, bild.height) <= 1.6 or min(bild.size) < 80:
                return None
            bild.thumbnail((FOTO_PX, FOTO_PX))
            if bild.mode in ("RGBA", "LA", "P"):
                bild = bild.convert("RGBA")
                weiss = Image.new("RGB", bild.size, "white")
                weiss.paste(bild, mask=bild.split()[-1])
                bild = weiss
            puffer = io.BytesIO()
            bild.convert("RGB").save(puffer, "JPEG", quality=88)
            return puffer.getvalue()
    except Exception as e:
        print("Foto unbrauchbar:", repr(e))
        return None


def foto(pfad):
    """Das Foto aus dem hochgeladenen Lebenslauf als JPEG-bytes oder None.
    Genommen wird das groesste Bild (Logos und Symbole sind kleiner)."""
    p = Path(pfad)
    kandidaten = []
    try:
        if p.suffix.lower() == ".pdf":
            from pypdf import PdfReader
            for seite in PdfReader(p).pages[:2]:
                for bild in seite.images:
                    try:                               # ein kaputtes Bild kostet nur sich
                        daten = bild.data
                        if len(daten) > MAX_BILD_BYTES:
                            continue
                        if not _bildart(daten):        # rohe PDF-Bilddaten
                            roh = bild.image
                            if roh.width * roh.height > MAX_PIXEL:
                                continue
                            puffer = io.BytesIO()
                            roh.convert("RGB").save(puffer, "JPEG", quality=90)
                            daten = puffer.getvalue()
                        kandidaten.append(daten)
                    except Exception as e:
                        print("PDF-Bild übersprungen:", repr(e))
        else:
            with zipfile.ZipFile(p) as z:
                kandidaten = [z.read(i) for i in z.infolist()
                              if i.filename.startswith("word/media/")
                              and MIN_FOTO <= i.file_size <= MAX_BILD_BYTES]
    except Exception as e:
        print("Foto nicht lesbar:", repr(e))
        return None
    kandidaten = sorted((k for k in kandidaten if _bildart(k) and len(k) >= MIN_FOTO),
                        key=len, reverse=True)
    for k in kandidaten[:3]:
        sauber = _sauber(k)
        if sauber:
            return sauber
    return None


def lebenslauf_text(pfad):
    p = Path(pfad)
    if p.suffix.lower() == ".pdf":
        from pypdf import PdfReader
        return "\n".join((s.extract_text() or "") for s in PdfReader(p).pages).strip()
    # Direkt aus dem XML: python-docx scheitert an manchen Word-Dateien
    # (z. B. Bilder ohne Dateityp).
    xml = zipfile.ZipFile(p).read("word/document.xml").decode("utf-8")
    absaetze = []
    for absatz in re.findall(r"<w:p[ >].*?</w:p>", xml, re.S):
        teile = re.findall(r"<w:t[^>]*>([^<]*)</w:t>|(<w:tab/>)", absatz)
        text = "".join(t or "\t" for t, _ in teile)
        if text.strip():
            absaetze.append(text)
    import html
    return html.unescape("\n".join(absaetze))


def _fett(absatz, text):
    for i, stueck in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if stueck:
            absatz.add_run(stueck).bold = (i % 2 == 1)


def lebenslauf_docx(markdown, ziel, bild=None):
    """bild: bytes des Fotos (png/jpg) oder None. Mit Foto stehen Name und
    Berufstitel links in einer randlosen Tabelle, das Foto rechts daneben."""
    doc = Document()
    stil = doc.styles["Normal"]
    stil.font.name, stil.font.size = "Calibri", Pt(10.5)
    zeilen = [z.rstrip() for z in markdown.splitlines() if z.strip()]
    if bild and not (zeilen and zeilen[0].startswith("# ")):
        # Ohne Namenszeile: Foto trotzdem oben rechts.
        absatz = doc.add_paragraph()
        absatz.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        absatz.add_run().add_picture(io.BytesIO(bild), width=Cm(3.5))
    elif bild:
        tabelle = doc.add_table(rows=1, cols=2)
        links, rechts = tabelle.rows[0].cells
        links.width, rechts.width = Cm(12.5), Cm(4)
        r = links.paragraphs[0].add_run(zeilen[0][2:].strip())
        r.bold, r.font.size = True, Pt(20)
        zeilen = zeilen[1:]
        if zeilen and not zeilen[0].startswith("#") and not re.match(r"^\s*[-•*] ", zeilen[0]):
            _fett(links.add_paragraph(), zeilen[0])
            zeilen = zeilen[1:]
        absatz = rechts.paragraphs[0]
        absatz.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        absatz.add_run().add_picture(io.BytesIO(bild), width=Cm(3.5))
    for z in zeilen:
        if z.startswith("# "):
            r = doc.add_paragraph().add_run(z[2:].strip())
            r.bold, r.font.size = True, Pt(20)
        elif z.startswith("## "):
            a = doc.add_paragraph()
            a.paragraph_format.space_before = Pt(10)
            r = a.add_run(z[3:].strip().upper())
            r.bold, r.font.size = True, Pt(11)
            r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x8C)
        elif re.match(r"^\s*[-•*] ", z):
            _fett(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-•*] ", "", z))
        else:
            _fett(doc.add_paragraph(), z.lstrip("# "))
    doc.save(ziel)


def anschreiben_docx(text, ziel):
    doc = Document()
    stil = doc.styles["Normal"]
    stil.font.name, stil.font.size = "Calibri", Pt(11)
    for block in re.split(r"\n\s*\n", text.strip()):
        a = doc.add_paragraph()
        zeilen = block.splitlines()
        for i, z in enumerate(zeilen):
            _fett(a, z.strip())
            if i < len(zeilen) - 1:
                a.add_run().add_break()
    doc.save(ziel)
